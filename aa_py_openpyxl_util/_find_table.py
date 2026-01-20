from logging import getLogger
from typing import Tuple, TYPE_CHECKING, Literal

from pydicti import dicti

if TYPE_CHECKING:
    from openpyxl import Workbook
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.table import Table
    from openpyxl.worksheet.worksheet import Worksheet


logger = getLogger(__name__)


def find_table(
    *,
    book: "Workbook",
    name: str,
    ci: (
        bool | Literal["warn"]
    ) = False,  # TODO: Make this required in the next major version.
) -> Tuple["Worksheet", str]:
    """
    Find a table in the given workbook. The table can be a named range or a ListObject.

    Args:
        book: The workbook to search through.
        name: The name of the table (named range or ListObject).
        ci:
            Whether the table name lookup should be case-insensitive.
            When this is "warn", a warning is logged when the provided case does not match the actual case.

    Returns:
        A tuple of (sheet, range).
    """
    try:
        sheet, table_range = find_named_range_by_name(book=book, name=name, ci=ci)
        table_type = "Named range"
    except KeyError as e1:
        try:
            sheet, table = find_list_object_by_name(book=book, name=name, ci=ci)
            table_range = table.ref
            table_type = "ListObject"
        except KeyError as e2:
            raise KeyError(f"Table `{name}` not found: {e1.args[0]} {e2.args[0]}")

    from openpyxl.utils import range_boundaries

    min_col, min_row, max_col, max_row = range_boundaries(table_range)
    n_rows = max_row - min_row + 1
    n_cols = max_col - min_col + 1

    if n_rows < 2:
        raise KeyError(f"{table_type} `{name}` found, but it has fewer than 2 rows.")
    if n_cols < 1:
        raise KeyError(f"{table_type} `{name}` found, but it has no columns.")

    return sheet, table_range


def find_named_range_by_name(
    *,
    book: "Workbook",
    name: str,
    ci: bool | Literal["warn"],
) -> Tuple["Worksheet", str]:
    """
    Find a named range by name.

    Args:
        book: The book to search through.
        name: The name of the named range.
        ci:
            Whether the table name lookup should be case-insensitive.
            When this is "warn", a warning is logged when the provided case does not match the actual case.

    Returns:
        A tuple of (sheet, range).
    """
    book_defined_names = dicti(book.defined_names) if ci else book.defined_names
    try:
        defined_name: "DefinedName" = book_defined_names[name]
    except KeyError as e:
        raise KeyError(f"Named range `{name}` not found.") from e

    if ci == "warn":
        # Check for case mismatch
        original_name = dicti((k, k) for k in book.defined_names.keys())[name]
        if original_name != name:
            logger.warning(
                f"Table with exact name `{name}` not found. Using case-insensitive match `{original_name}` instead."
            )

    try:
        destinations = list(defined_name.destinations)
    except AttributeError:
        raise KeyError(f"Named range `{name}` found, but it has no destinations.")

    if len(destinations) != 1:
        raise KeyError(f"Named range `{name}` found, but it has multiple destinations.")

    sheet_name, table_range = destinations[0]
    return book[sheet_name], table_range


def find_list_object_by_name(
    *,
    book: "Workbook",
    name: str,
    ci: bool | Literal["warn"],
) -> Tuple["Worksheet", "Table"]:
    """
    Find a list object by name.

    Args:
        book: The book to search through.
        name: The name of the ListObject.
        ci:
            Whether the table name lookup should be case-insensitive.
            When this is "warn", a warning is logged when the provided case does not match the actual case.

    Returns:
        A tuple of (sheet, table).
    """
    sheet: "Worksheet"
    for sheet in book.worksheets:
        sheet_tables = dicti(sheet.tables) if ci else sheet.tables
        if name in sheet_tables:
            if ci == "warn":
                # Check for case mismatch
                original_name = dicti((k, k) for k in sheet.tables.keys())[name]
                if original_name != name:
                    logger.warning(
                        f"Table with exact name `{name}` not found. Using case-insensitive match `{original_name}` instead."
                    )

            return sheet, sheet_tables[name]

    raise KeyError(f"ListObject `{name}` not found.")
